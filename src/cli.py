"""CLI offline: python -m src.cli --text 'địa chỉ' hoặc --input addresses.csv."""

import argparse
import csv
from dataclasses import asdict, fields
import json
from pathlib import Path
import sys
from .chuan_hoa import BoChuanHoa
from .kieu_du_lieu import KetQua


def read_rows(path):
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as e:
            raise RuntimeError("Cài requirements.txt để đọc Excel.") from e
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            header = next(rows, ())
            return [dict(zip(header, row)) for row in rows]
        finally:
            workbook.close()
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_rows(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    header = list(rows[0]) if rows else ["dau_vao"] + [f.name for f in fields(KetQua)]
    serialized = [
        {
            k: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v
            for k, v in r.items()
        }
        for r in rows
    ]
    if path.suffix.lower() == ".xlsx":
        import openpyxl

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(header)
        for r in serialized:
            sheet.append([r.get(k, "") for k in header])
        # Văn bản nhập vào, kể cả dấu '=' đầu chuỗi, phải là Text thay vì công thức Excel.
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        book.save(path)
    elif path.suffix.lower() == ".json":
        path.write_text(
            json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    else:
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=header)
            writer.writeheader()
            writer.writerows(serialized)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description=__doc__)
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--text")
    g.add_argument("--input", type=Path)
    p.add_argument("--output", type=Path)
    p.add_argument("--column", default="dau_vao")
    p.add_argument("--no-fuzzy", action="store_true")
    args = p.parse_args()
    model = BoChuanHoa(fuzzy=not args.no_fuzzy)
    if args.text is not None:
        result = asdict(model.chuan_hoa(args.text))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if args.output:
            write_rows(args.output, [dict(dau_vao=args.text, **result)])
        return
    if not args.output:
        p.error("--input cần --output để không ghi đè dữ liệu gốc.")
    if args.input.resolve() == args.output.resolve():
        p.error("Tệp đầu ra phải khác tệp đầu vào.")
    rows = read_rows(args.input)
    results = []
    if rows and args.column not in rows[0]:
        p.error(f"Không có cột {args.column!r}.")
    for r in rows:
        text = r.get(args.column) or ""
        try:
            result = asdict(model.chuan_hoa(str(text)))
        except (ValueError, TypeError) as e:
            result = asdict(KetQua(canh_bao=["dau_vao_khong_hop_le", str(e)]))
        results.append(dict(r, **result))
    write_rows(args.output, results)
    print(f"Đã xử lý {len(results)} dòng → {args.output}")


if __name__ == "__main__":
    main()
